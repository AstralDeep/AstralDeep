"""Tests for the opt-in structured PDF/DOCX reader: ordering, tables, provenance, bounded
worker containment, fail-closed assets, denial, cancellation, timeout, and malformed input.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import subprocess
import sys
import threading
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

from agents.general.file_tools import structured_document
from agents.general.file_tools.read_document import read_document
from agents.general.file_tools.structured_document import parse_structured_document
from conftest import _persist
from shared.feature_flags import flags


def _enable(monkeypatch):
    monkeypatch.setitem(flags._flags, "structured_document_parser", True)


def _pdf(pages, *, font_size=10):
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    for operators in pages:
        document.setFont("Helvetica", font_size)
        for x, y, text in operators:
            document.drawString(x, y, text)
        document.showPage()
    document.save()
    return buffer.getvalue()


def _table_pdf(rows, *, x=80, y=740, step=20, column=120):
    return _pdf([[
        (x + column * index, y - step * row, value)
        for row, values in enumerate(rows)
        for index, value in enumerate(values)
    ]])


def _columns_pdf():
    rows = [
        ("Left column line one with several words", "Right column line one with several words"),
        ("Left column line two with several words", "Right column line two with several words"),
        ("Left column line three with several words", "Right column line three with several words"),
    ]
    return _pdf([[
        item
        for index, (left, right) in enumerate(rows)
        for item in ((60, 740 - index * 18, left), (330, 740 - index * 18, right))
    ]], font_size=12)


def _scanned_pdf():
    from PIL import Image
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    image = Image.new("RGB", (200, 100), "white")
    payload = io.BytesIO()
    image.save(payload, "PNG")
    payload.seek(0)
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawImage(ImageReader(payload), 100, 600, width=200, height=100)
    document.showPage()
    document.save()
    return buffer.getvalue()


def _encrypted_pdf():
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(_pdf([[(100, 750, "secret text that must not be read")]])))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.encrypt(user_password="user", owner_password="owner")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx(blocks):
    import docx

    document = docx.Document()
    for kind, value in blocks:
        if kind == "p":
            document.add_paragraph(value)
        else:
            table = document.add_table(rows=len(value), cols=len(value[0]))
            for row_index, row in enumerate(value):
                for column_index, cell in enumerate(row):
                    table.rows[row_index].cells[column_index].text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_table_payload():
    return _docx([
        ("p", "Before table."),
        ("table", [["Analyte", "Value", "Unit"], ["Sodium", "-1.25", "mg"]]),
        ("p", "After table."),
    ])


def _rewrite_zip(payload, changes):
    source = zipfile.ZipFile(io.BytesIO(payload))
    output = io.BytesIO()
    written = set()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = changes.get(info.filename, source.read(info.filename))
            target.writestr(info, data)
            written.add(info.filename)
        for name, data in changes.items():
            if name not in written:
                target.writestr(name, data)
    return output.getvalue()


def _docx_with_external_relationship(payload):
    relation_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        root = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
    ElementTree.SubElement(root, f"{{{relation_ns}}}Relationship", {
        "Id": "rId999",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        "Target": "https://example.invalid/image.png",
        "TargetMode": "External",
    })
    return _rewrite_zip(payload, {
        "word/_rels/document.xml.rels": ElementTree.tostring(root),
    })


def _profile_variant(tmp_path, monkeypatch, **limits):
    profile = json.loads(
        (Path(__file__).resolve().parents[5] / "config" / "documents" /
         "structured-reader-profile.json").read_text()
    )
    profile["limits"].update(limits)
    path = tmp_path / "profile.json"
    path.write_text(json.dumps(profile))
    monkeypatch.setenv("ASTRAL_STRUCTURED_DOCUMENT_PROFILE", str(path))
    return path


def _wait_for_file(path, timeout=3.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if path.exists():
            return path.read_text()
        time.sleep(0.02)
    raise AssertionError("worker helper did not record its child")


def _assert_pid_gone(pid, timeout=3.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            os.kill(pid, 0)
        except ProcessLookupError:
            return
        time.sleep(0.05)
    raise AssertionError("supervised worker descendant survived cleanup")


def test_flag_defaults_off_and_structured_request_fails_closed(repo, upload_root, monkeypatch):
    monkeypatch.setitem(flags._flags, "structured_document_parser", False)
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "A plain document with enough readable text.")]]))
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "failed"
    assert out["error"]["code"] == "structured_parser_disabled"
    assert "text" not in out


def test_structured_pdf_returns_blocks_regions_and_source_identity(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    payload = _pdf([[(100, 750, "Structured text with a genuine source region.")]])
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=payload)
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "complete" and out["complete"] is True
    assert out["parser"]["adapter"] == "structured-pdf-docx"
    assert out["source"]["sha256"] == hashlib.sha256(payload).hexdigest()
    assert out["blocks"][0]["source"]["page"] == 1
    region = out["blocks"][0]["source"]["region"]
    assert len(region) == 4 and region[0] < region[2] and region[1] < region[3]
    assert out["worker"]["cleanup_ok"] is True


def test_structured_pdf_multicolumn_reading_order(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    payload = _columns_pdf()
    aid = _persist(repo, user_id="alice", filename="columns.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=payload)
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "complete"
    assert out["tables"] == []
    text = out["text"]
    assert text.index("Left column line three") < text.index("Right column line one")


def test_structured_pdf_table_preserves_cells_and_numeric_text(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    rows = [["Analyte", "Value", "Unit"], ["Sodium", "-1.25", "mg"],
            ["Glucose", "+2.50", "mmol/L"]]
    aid = _persist(repo, user_id="alice", filename="table.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_table_pdf(rows))
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "complete"
    table = out["tables"][0]
    assert [cell["text"] for cell in table["rows"][1]["cells"]] == ["Sodium", "-1.25", "mg"]
    for row in table["rows"]:
        for cell in row["cells"]:
            assert cell["source"]["page"] == 1
            assert len(cell["source"]["region"]) == 4
    assert "-1.25" in out["text"] and "+2.50" in out["text"]


def test_structured_pdf_selected_page_keeps_original_identity(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    payload = _pdf([
        [(100, 750, "First page text that should not be selected.")],
        [(100, 750, "Second page text that must keep page identity two.")],
    ])
    aid = _persist(repo, user_id="alice", filename="pages.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=payload)
    out = read_document(attachment_id=aid, user_id="alice", page_range="2", structured=True)
    assert out["selected_pages"] == [2]
    assert "Second page text" in out["text"] and "First page text" not in out["text"]
    assert {block["source"]["page"] for block in out["blocks"]} == {2}


def test_structured_pdf_blank_is_complete_empty_but_scanned_is_partial(
    repo, upload_root, monkeypatch
):
    _enable(monkeypatch)
    blank = _persist(repo, user_id="alice", filename="blank.pdf", category="document",
                     extension="pdf", content_type="application/pdf", upload_root=upload_root,
                     payload=_pdf([[]]))
    blank_out = read_document(attachment_id=blank, user_id="alice", structured=True)
    assert blank_out["status"] == "complete" and blank_out["text"] == ""
    assert blank_out["vision_required"] is False

    scanned = _persist(repo, user_id="alice", filename="scan.pdf", category="document",
                       extension="pdf", content_type="application/pdf", upload_root=upload_root,
                       payload=_scanned_pdf())
    scanned_out = read_document(attachment_id=scanned, user_id="alice", structured=True)
    assert scanned_out["status"] == "partial" and scanned_out["complete"] is False
    assert scanned_out["vision_required"] is True
    assert "page_image_without_text" in scanned_out["warnings"]
    assert scanned_out["text"] == "" and scanned_out["images"] == []


@pytest.mark.parametrize(
    ("payload", "code"),
    [(_encrypted_pdf(), "encrypted_document"),
     (b"%PDF-1.4\nnot a valid pdf body \x00\x01", "malformed_document")],
)
def test_structured_pdf_encrypted_and_malformed_fail_closed(
    repo, upload_root, monkeypatch, payload, code
):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="broken.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=payload)
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "failed" and out["complete"] is False
    assert out["error"]["code"] == code
    assert out["source"]["sha256"] == hashlib.sha256(payload).hexdigest()


def test_structured_docx_tables_have_block_identifiers_without_page_coordinates(
    repo, upload_root, monkeypatch
):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="memo.docx", category="document",
                   extension="docx",
                   content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                   upload_root=upload_root, payload=_docx_table_payload())
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "complete"
    assert [block["type"] for block in out["blocks"]] == ["paragraph", "table", "paragraph"]
    table = out["tables"][0]
    assert table["source"]["block_id"] == "body/tbl[0]"
    assert table["rows"][1]["cells"][1]["source"]["block_id"] == "body/tbl[0]/tr[1]/tc[1]"
    assert table["rows"][1]["cells"][1]["text"] == "-1.25"
    assert "page" not in json.dumps(table)
    assert "After table." in out["text"]


def test_structured_docx_marks_continuation_tables(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    payload = _docx([
        ("table", [["Item", "Value"], ["First", "1.5"]]),
        ("table", [["Item", "Value"], ["Second", "-2.50"]]),
    ])
    aid = _persist(repo, user_id="alice", filename="continued.docx", category="document",
                   extension="docx",
                   content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                   upload_root=upload_root, payload=payload)
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["tables"][1]["continuation_of"] == "body/tbl[0]"
    assert out["tables"][1]["header_repeated"] is True


def test_structured_docx_external_relationship_and_macro_fail_closed(
    repo, upload_root, monkeypatch
):
    _enable(monkeypatch)
    cases = [
        (_docx_with_external_relationship(_docx_table_payload()), "external_resource_rejected"),
        (_rewrite_zip(_docx_table_payload(), {"word/vbaProject.bin": b"macro"}),
         "active_content_rejected"),
        (b"PK\x03\x04broken archive bytes", "malformed_document"),
    ]
    for index, (payload, code) in enumerate(cases):
        aid = _persist(repo, user_id="alice", filename=f"case{index}.docx",
                       category="document", extension="docx",
                       content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                       upload_root=upload_root, payload=payload)
        out = read_document(attachment_id=aid, user_id="alice", structured=True)
        assert out.get("status") == "failed", out
        assert out["error"]["code"] == code
        assert out["text"] == ""


def test_structured_output_truncation_is_truthful(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="long.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "x" * 500)]]))
    out = read_document(attachment_id=aid, user_id="alice", max_chars=20, structured=True)
    assert out["status"] == "truncated" and out["complete"] is False
    assert out["truncated"] is True
    assert len(out["text"]) <= 20


def test_structured_missing_profile_dependency_drift_and_worker_fail_closed(
    repo, upload_root, monkeypatch, tmp_path
):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "Text for availability checks.")]]))

    monkeypatch.setenv("ASTRAL_STRUCTURED_DOCUMENT_PROFILE", str(tmp_path / "missing.json"))
    missing = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert missing["error"]["code"] == "structured_parser_unavailable"

    profile_path = _profile_variant(tmp_path, monkeypatch)
    profile = json.loads(profile_path.read_text())
    profile["dependencies"][0]["version"] = "0.0.0"
    profile_path.write_text(json.dumps(profile))
    drift = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert drift["status"] == "failed"
    assert drift["error"]["code"] == "dependency_mismatch"

    monkeypatch.delenv("ASTRAL_STRUCTURED_DOCUMENT_PROFILE")
    monkeypatch.setattr(structured_document, "WORKER_PATH", tmp_path / "missing_worker.py")
    no_worker = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert no_worker["error"]["code"] == "structured_parser_unavailable"


def test_structured_owner_denial_happens_before_worker(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "Owned text.")]]))
    out = read_document(attachment_id=aid, user_id="mallory", structured=True)
    assert out["error"]["code"] == "not_found"
    assert "text" not in out


def test_structured_cancellation_terminates_the_worker_and_descendants(
    monkeypatch, tmp_path
):
    _profile_variant(tmp_path, monkeypatch, processes=1024)
    pid_file = tmp_path / "cancel-child.pid"
    worker = tmp_path / "slow_worker.py"
    worker.write_text(
        "import subprocess, sys, time\n"
        f"child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\n"
        f"open({str(pid_file)!r}, 'w').write(str(child.pid))\n"
        "time.sleep(60)\n"
    )
    monkeypatch.setattr(structured_document, "WORKER_PATH", worker)
    event = threading.Event()
    timer = threading.Timer(0.3, event.set)
    timer.start()
    try:
        out = parse_structured_document(
            _pdf([[(100, 750, "Cancel me.")]]), "pdf", None, 200000,
            cancel_event=event,
        )
    finally:
        timer.cancel()
    assert out.get("status") == "cancelled" and out.get("cleanup_ok") is True, out
    _assert_pid_gone(int(_wait_for_file(pid_file)))


def test_structured_timeout_terminates_the_worker_and_descendants(
    repo, upload_root, monkeypatch, tmp_path
):
    _enable(monkeypatch)
    _profile_variant(tmp_path, monkeypatch, deadline_seconds=0.5, processes=1024)
    pid_file = tmp_path / "timeout-child.pid"
    worker = tmp_path / "timeout_worker.py"
    worker.write_text(
        "import subprocess, sys, time\n"
        f"child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\n"
        f"open({str(pid_file)!r}, 'w').write(str(child.pid))\n"
        "time.sleep(60)\n"
    )
    monkeypatch.setattr(structured_document, "WORKER_PATH", worker)
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "Time out.")]]))
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "timeout" and out["cleanup_ok"] is True
    _assert_pid_gone(int(_wait_for_file(pid_file)))


@pytest.mark.parametrize(
    ("script", "limits"),
    [
        ("from pathlib import Path\nPath('output.json').write_text('not json')\n", {}),
        ("raise SystemExit(7)\n", {}),
        ("from pathlib import Path\nPath('output.json').write_text('{\"x\":\"' + 'a' * 4096 + '\"}')\n",
         {"result_bytes": 1024}),
    ],
)
def test_structured_invalid_or_oversized_worker_output_fails_closed(
    repo, upload_root, monkeypatch, tmp_path, script, limits
):
    _enable(monkeypatch)
    if limits:
        _profile_variant(tmp_path, monkeypatch, **limits)
    worker = tmp_path / "broken_worker.py"
    worker.write_text(script)
    monkeypatch.setattr(structured_document, "WORKER_PATH", worker)
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "Broken worker.")]]))
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["error"]["code"] == "structured_parser_failed"
    assert "text" not in out


def test_structured_input_overflow_and_page_limit_fail_or_truncate_truthfully(
    repo, upload_root, monkeypatch, tmp_path
):
    _enable(monkeypatch)
    payload = _pdf([
        [(100, 750, "Page one text.")],
        [(100, 750, "Page two text.")],
    ])
    aid = _persist(repo, user_id="alice", filename="pages.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=payload)

    _profile_variant(tmp_path, monkeypatch, input_bytes=100)
    overflow = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert overflow["status"] == "failed"
    assert overflow["error"]["code"] == "structured_input_too_large"

    _profile_variant(tmp_path, monkeypatch, pdf_pages=1)
    limited = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert limited["status"] == "truncated" and limited["complete"] is False
    assert limited["selected_pages"] == [1]
    assert "page_limit_truncated" in limited["warnings"]


def test_structured_docx_archive_limit_fails_closed(repo, upload_root, monkeypatch, tmp_path):
    _enable(monkeypatch)
    _profile_variant(tmp_path, monkeypatch, docx_archive_entries=1)
    aid = _persist(repo, user_id="alice", filename="memo.docx", category="document",
                   extension="docx",
                   content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                   upload_root=upload_root, payload=_docx_table_payload())
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "failed" and out["error"]["code"] == "archive_limit"


def test_baseline_reader_is_unchanged_when_flag_on_but_structured_not_requested(
    repo, upload_root, monkeypatch
):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="memo.pdf", category="document",
                   extension="pdf", content_type="application/pdf", upload_root=upload_root,
                   payload=_pdf([[(100, 750, "Baseline text remains the default path.")]]))
    out = read_document(attachment_id=aid, user_id="alice")
    assert "status" not in out and "parser" not in out
    assert "Baseline text remains" in out["text"]


def test_structured_rejects_unsupported_extension(repo, upload_root, monkeypatch):
    _enable(monkeypatch)
    aid = _persist(repo, user_id="alice", filename="memo.rtf", category="document",
                   extension="rtf", content_type="application/rtf", upload_root=upload_root,
                   payload=b"{\\rtf1\\ansi text}")
    out = read_document(attachment_id=aid, user_id="alice", structured=True)
    assert out["status"] == "failed" and out["error"]["code"] == "unsupported"


def test_worker_network_guard_denies_outbound_connections():
    backend = Path(__file__).resolve().parents[4]
    code = (
        "from agents.general.file_tools.structured_worker import _install_network_guard\n"
        "_install_network_guard()\n"
        "import socket\n"
        "try:\n"
        "    socket.create_connection(('127.0.0.1', 9), timeout=0.2)\n"
        "except OSError:\n"
        "    raise SystemExit(0)\n"
        "raise SystemExit(1)\n"
    )
    result = subprocess.run(
        [sys.executable, "-c", code],
        cwd=backend.parent,
        env={**os.environ, "PYTHONPATH": str(backend)},
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 0