"""In-process tests for the structured worker's parser, profile, preflight, request, and
result boundaries; production execution still goes through the supervised child process.
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest

from agents.general.file_tools import structured_worker as worker


def _profile():
    return json.loads(
        (Path(__file__).resolve().parents[5] / "config" / "documents" /
         "structured-reader-profile.json").read_text()
    )


def _pdf(pages):
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    for operators in pages:
        document.setFont("Helvetica", 10)
        for x, y, text in operators:
            document.drawString(x, y, text)
        document.showPage()
    document.save()
    return buffer.getvalue()


def _docx(paragraphs=(), tables=()):
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.rows[row_index].cells[column_index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _request(tmp_path, monkeypatch, payload, extension="pdf", profile=None,
             cancel_guard=True):
    profile = profile or _profile()
    runtime = tmp_path / "runtime"
    work = tmp_path / "work"
    runtime.mkdir(parents=True, exist_ok=True)
    work.mkdir(parents=True, exist_ok=True)
    profile_raw = json.dumps(profile).encode()
    (runtime / "profile.json").write_bytes(profile_raw)
    (work / "input.bin").write_bytes(payload)
    request = {
        "version": 1,
        "extension": extension,
        "page_range": None,
        "max_chars": 200_000,
        "input": "input.bin",
        "output": "output.json",
        "profile": "profile.json",
        "profile_sha256": hashlib.sha256(profile_raw).hexdigest(),
    }
    request_path = work / "request.json"
    request_path.write_text(json.dumps(request))
    monkeypatch.chdir(work)
    if cancel_guard:
        monkeypatch.setattr(worker, "_install_network_guard", lambda: None)
    return request_path, work / "output.json", runtime / "profile.json"


def test_run_parses_pdf_and_docx_in_process(tmp_path, monkeypatch):
    request, output, _profile_path = _request(
        tmp_path, monkeypatch, _pdf([[(100, 750, "In-process worker text.")]])
    )
    result = worker._run(request)
    assert result["status"] == "complete"
    assert json.loads(output.read_text())["source"]["extension"] == "pdf"

    docx_request, docx_output, _ = _request(
        tmp_path / "docx", monkeypatch,
        _docx(["Before."], [[["Item", "Value"], ["Sodium", "-1.25"]]]),
        extension="docx",
    )
    docx_result = worker._run(docx_request)
    assert docx_result["tables"][0]["rows"][1]["cells"][1]["text"] == "-1.25"
    assert json.loads(docx_output.read_text())["tables"]


def test_main_writes_a_bounded_failure_for_profile_digest_mismatch(tmp_path, monkeypatch):
    request, output, _profile_path = _request(
        tmp_path, monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    value = json.loads(request.read_text())
    value["profile_sha256"] = "0" * 64
    request.write_text(json.dumps(value))
    assert worker.main(["worker", str(request)]) == 0
    assert json.loads(output.read_text())["error"]["code"] == "profile_digest_mismatch"


@pytest.mark.parametrize(
    ("mutate", "code"),
    [
        (lambda value: value.update(version=2), "invalid_request"),
        (lambda value: value.update(extension="rtf"), "unsupported_extension"),
        (lambda value: value.update(max_chars=0), "invalid_request"),
        (lambda value: value.update(input="../escape.bin"), "invalid_request"),
        (lambda value: value.update(profile_sha256="0" * 64), "profile_digest_mismatch"),
    ],
)
def test_run_rejects_invalid_requests(tmp_path, monkeypatch, mutate, code):
    request, _output, _profile_path = _request(
        tmp_path, monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    value = json.loads(request.read_text())
    mutate(value)
    request.write_text(json.dumps(value))
    with pytest.raises(worker.WorkerFailure) as error:
        worker._run(request)
    assert error.value.code == code


def test_run_rejects_dependency_drift_input_overflow_platform_and_root(tmp_path, monkeypatch):
    profile = _profile()
    profile["dependencies"][0]["version"] = "0.0.0"
    request, _output, _profile_path = _request(
        tmp_path / "drift", monkeypatch, _pdf([[(100, 750, "Text.")]]), profile=profile
    )
    with pytest.raises(worker.WorkerFailure, match="dependency_mismatch"):
        worker._run(request)

    profile = _profile()
    profile["limits"]["input_bytes"] = 10
    request, _output, _profile_path = _request(
        tmp_path / "overflow", monkeypatch, _pdf([[(100, 750, "Text.")]]), profile=profile
    )
    with pytest.raises(worker.WorkerFailure, match="input_too_large"):
        worker._run(request)

    request, _output, _profile_path = _request(
        tmp_path / "platform", monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    monkeypatch.setattr(worker.sys, "platform", "darwin")
    with pytest.raises(worker.WorkerFailure, match="unsupported_platform"):
        worker._run(request)

    request, _output, _profile_path = _request(
        tmp_path / "root", monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    monkeypatch.setattr(worker.sys, "platform", "linux")
    monkeypatch.setattr(worker.os, "geteuid", lambda: 0)
    with pytest.raises(worker.WorkerFailure, match="root_worker_rejected"):
        worker._run(request)


@pytest.mark.parametrize(
    "mutate",
    [
        lambda profile: profile.update(schema_version=2),
        lambda profile: profile.pop("profile_id"),
        lambda profile: profile["adapter"].update(version=""),
        lambda profile: profile["limits"].pop("input_bytes"),
        lambda profile: profile["limits"].update(input_bytes=0),
        lambda profile: profile["limits"].update(input_bytes=10**12),
        lambda profile: profile["limits"].update(deadline_seconds=0),
        lambda profile: profile["limits"].update(deadline_seconds="20"),
        lambda profile: profile["containment"].update(network="allowed"),
        lambda profile: profile["containment"].update(environment="inherited"),
        lambda profile: profile["enrichment"].update(ocr=True),
        lambda profile: profile.update(assets=[{"path": "model.bin"}]),
        lambda profile: profile.update(dependencies=[{"name": "x"}]),
    ],
)
def test_validate_profile_rejects_boundary_violations(mutate):
    profile = _profile()
    mutate(profile)
    with pytest.raises(worker.WorkerFailure, match="invalid_profile"):
        worker._validate_profile(profile)


def test_page_range_parsing_preserves_order_and_bounds():
    assert worker._parse_page_range("3,1-2,2", 3, 10) == ([3, 1, 2], False)
    assert worker._parse_page_range("", 2, 1) == ([1], True)
    with pytest.raises(worker.WorkerFailure, match="invalid_page_range"):
        worker._parse_page_range("3-1", 3, 10)
    with pytest.raises(worker.WorkerFailure, match="invalid_page_range"):
        worker._parse_page_range("a", 3, 10)
    with pytest.raises(worker.WorkerFailure, match="empty_document"):
        worker._parse_page_range("", 0, 10)


def test_docx_preflight_rejects_archive_limits_paths_and_active_content():
    def archive(entries, compression=zipfile.ZIP_STORED):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression) as target:
            for name, data in entries.items():
                target.writestr(name, data)
        return buffer.getvalue()

    limits = _profile()["limits"]
    with pytest.raises(worker.WorkerFailure, match="malformed_document"):
        worker._preflight_docx(archive({"[Content_Types].xml": b"<Types/>"}), limits)
    with pytest.raises(worker.WorkerFailure, match="unsafe_archive"):
        worker._preflight_docx(
            archive({"word/document.xml": b"<x/>", "../escape": b"x"}), limits
        )
    with pytest.raises(worker.WorkerFailure, match="active_content_rejected"):
        worker._preflight_docx(
            archive({"word/document.xml": b"<x/>", "word/vbaProject.bin": b"x"}), limits
        )
    with pytest.raises(worker.WorkerFailure, match="external_resource_rejected"):
        worker._preflight_docx(
            archive({
                "word/document.xml": b"<x/>",
                "word/_rels/document.xml.rels": (
                    b'<Relationships xmlns="http://schemas.openxmlformats.org/'
                    b'package/2006/relationships"><Relationship TargetMode="External" '
                    b'Target="https://example.invalid"/></Relationships>'
                ),
            }),
            limits,
        )
    constrained = dict(limits)
    constrained["docx_archive_entries"] = 1
    with pytest.raises(worker.WorkerFailure, match="archive_limit"):
        worker._preflight_docx(
            archive({"word/document.xml": b"<x/>", "word/styles.xml": b"<x/>"}),
            constrained,
        )
    constrained = dict(limits)
    constrained["docx_entry_bytes"] = 4
    with pytest.raises(worker.WorkerFailure, match="archive_limit"):
        worker._preflight_docx(archive({"word/document.xml": b"x" * 5}), constrained)


def test_docx_nested_table_has_path_without_duplicating_cell_paragraphs():
    import docx

    document = docx.Document()
    outer = document.add_table(rows=1, cols=1)
    cell = outer.rows[0].cells[0]
    cell.text = "Outer cell paragraph."
    nested = cell.add_table(rows=1, cols=1)
    nested.rows[0].cells[0].text = "Nested value."
    buffer = io.BytesIO()
    document.save(buffer)
    result = worker._parse_docx(buffer.getvalue(), 200_000, _profile()["limits"], {})
    assert result["status"] == "complete"
    table_blocks = [block for block in result["blocks"] if block["type"] == "table"]
    assert len(table_blocks) == 2
    assert table_blocks[1]["source"]["block_id"].startswith("body/tbl[0]/tr[0]/tc[0]/tbl[0]")
    assert result["text"].count("Outer cell paragraph.") == 1


def test_write_result_shrinks_oversized_results_and_network_denial():
    import tempfile

    result = {
        "status": "complete",
        "complete": True,
        "truncated": False,
        "text": "x" * 5000,
        "blocks": [{"text": "x" * 5000}],
        "tables": [],
        "warnings": [],
    }
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "output.json"
        worker._write_result(path, result, 1024)
        value = json.loads(path.read_text())
        assert value["status"] == "truncated"
        assert value["truncated"] is True and value["complete"] is False
        assert len(path.read_bytes()) <= 1024
    with pytest.raises(OSError, match="network access denied"):
        worker._network_denied()


def test_safe_identity_and_clean_text_helpers():
    assert worker._safe_name("input.bin") == "input.bin"
    with pytest.raises(worker.WorkerFailure, match="invalid_request"):
        worker._safe_name("../input.bin")
    assert worker._clean_inline("  a\x00  b  ", 100) == "a b"
    assert worker._clean_block("a\r\n b\n\nc", 100) == "a\nb\nc"
    assert worker._bbox([10, 20, 0, 5]) == [0.0, 5.0, 10.0, 20.0]
    assert worker._merge_bbox([0, 0, 1, 1], [2, 2, 3, 3]) == [0.0, 0.0, 3.0, 3.0]


def test_table_likelihood_separates_short_cells_from_text_columns():
    short = [
        {"cells": [{"text": "Item"}, {"text": "Value"}]},
        {"cells": [{"text": "Sodium"}, {"text": "-1.25"}]},
    ]
    long_text = [
        {"cells": [{"text": "Left column sentence with many words"},
                   {"text": "Right column sentence with many words"}]},
        {"cells": [{"text": "Another left sentence with many words"},
                   {"text": "Another right sentence with many words"}]},
    ]
    assert worker._looks_like_table(short) is True
    assert worker._looks_like_table(long_text) is False


def test_pdf_helpers_cover_layout_and_image_detection():
    from pypdf import PdfReader

    payload = _pdf([[
        (60, 740, "Left line"), (330, 740, "Right line"),
        (60, 722, "Left second"), (330, 722, "Right second"),
    ]])
    page = PdfReader(io.BytesIO(payload)).pages[0]
    fragments = worker._pdf_fragments(page, _profile()["limits"])
    lines = worker._group_lines(fragments, _profile()["limits"])
    for line in lines:
        line["page"] = 1
    assert len(lines) == 2
    assert worker._page_has_image(page) is False
    assert worker._column_order(worker._split_cells(lines)) is not None
    assert worker._split_cells(lines)[0]["page"] if "page" in lines[0] else True


def test_run_missing_input_and_worker_failure_main_exit_paths(tmp_path, monkeypatch):
    request, output, _profile_path = _request(
        tmp_path, monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    (request.parent / "input.bin").unlink()
    with pytest.raises(worker.WorkerFailure, match="io_failure"):
        worker._run(request)
    assert worker.main([]) == 2
    assert worker.main(["worker", str(tmp_path / "missing-request.json")]) == 3

def _line(identifier, x, y, cells, page=1):
    cell_values = [
        {"text": text, "bbox": [x + index * 20, y - 5, x + index * 20 + 15, y + 5]}
        for index, text in enumerate(cells)
    ]
    return {
        "id": identifier,
        "y": y,
        "size": 10.0,
        "page": page,
        "cells": cell_values,
        "bbox": [x, y - 5, x + 20 * len(cells), y + 5],
        "text": " ".join(cells),
    }


def test_group_lines_merges_small_gaps_and_skips_bad_fragments():
    fragments = [
        {"text": "A", "x": 10, "y": 100, "x0": 10, "x1": 16, "y0": 98, "y1": 102, "size": 10},
        {"text": "B", "x": 18, "y": 100, "x0": 18, "x1": 24, "y0": 98, "y1": 102, "size": 10},
    ]
    lines = worker._group_lines(fragments, _profile()["limits"])
    assert lines[0]["text"] == "A B"

    class FakePage:
        def extract_text(self, visitor_text):
            visitor_text("", None, None, None, None)
            visitor_text("bad", [1, 0, 0, 1, 0, None], [1, 0, 0, 1, 0, 0], None, 10)
            visitor_text("nan", [1, 0, 0, 1, 0, 0], [1, 0, 0, 1, float("nan"), 0], None, 10)
            visitor_text("x\n", [1, 0, 0, 1, 0, 0], [1, 0, 0, 1, 10, 20], None, 10)

    assert worker._pdf_fragments(FakePage(), _profile()["limits"])[0]["text"] == "x"


def test_columns_match_table_detection_and_column_crossing_boundaries():
    left = _line("l", 10, 100, ["a", "b"])
    shifted = _line("s", 60, 90, ["c", "d"])
    assert worker._columns_match(left["cells"], left["cells"]) is True
    assert worker._columns_match(left["cells"], shifted["cells"]) is False
    assert worker._columns_match(left["cells"], left["cells"][:1]) is False
    assert worker._looks_like_table([left]) is False
    assert worker._looks_like_table([_line("t", 10, 100, ["a", "b"])]) is False

    runs = worker._detect_tables([
        _line("a", 10, 100, ["Item", "Value"]),
        _line("b", 10, 90, ["Sodium", "-1.25"]),
        _line("c", 10, 80, ["plain text"]),
        _line("d", 10, 70, ["Item", "Value"]),
        _line("e", 10, 60, ["Other", "2.50"]),
    ])
    assert len(runs) == 2

    crossing = [
        _line("l1", 10, 100, ["left"]),
        _line("l2", 10, 90, ["left"]),
        _line("r1", 200, 100, ["right"]),
        _line("r2", 200, 90, ["right"]),
        {"id": "wide", "y": 80, "size": 10, "page": 1,
         "cells": [{"text": "wide", "bbox": [10, 75, 250, 85]}],
         "bbox": [10, 75, 250, 85], "text": "wide"},
    ]
    assert worker._column_order(crossing) is None


def test_table_block_and_render_entry_limits():
    lines = [
        _line("a", 10, 100, ["Item", "Value"]),
        _line("b", 10, 90, ["Sodium", "-1.25"]),
    ]
    limits = _profile()["limits"]
    result = {"truncated": False, "warnings": [], "blocks": [], "tables": []}
    table = worker._table_block(1, 1, lines, {**limits, "table_rows": 1}, result)
    assert result["truncated"] is True and len(table["rows"]) == 1
    result = {"truncated": False, "warnings": [], "blocks": [], "tables": []}
    table = worker._table_block(1, 1, lines, {**limits, "table_cells": 1}, result)
    assert result["truncated"] is True and len(table["rows"][1]["cells"]) == 0

    entry = {"kind": "line", "item": lines[0]}
    result = {"truncated": False, "warnings": [], "blocks": [], "tables": []}
    text, exhausted = worker._render_entry(entry, 2, result, limits)
    assert text == "It" and exhausted is True and result["blocks"][0]["truncated"] is True
    result = {"truncated": False, "warnings": [], "blocks": [{}] * limits["blocks"], "tables": []}
    text, exhausted = worker._render_entry(entry, 100, result, {**limits, "blocks": 1})
    assert text == "" and exhausted is True

    table_entry = {"kind": "table", "item": table}
    result = {"truncated": False, "warnings": [], "blocks": [], "tables": []}
    text, exhausted = worker._render_entry(table_entry, 3, result, limits)
    assert text == "Ite" and exhausted is True and result["tables"][0]["truncated"] is True
    result = {"truncated": False, "warnings": [], "blocks": [{}] * limits["blocks"], "tables": []}
    text, exhausted = worker._render_entry(table_entry, 100, result, {**limits, "blocks": 1})
    assert text == "" and exhausted is True


def test_page_has_image_variants():
    class Page:
        def __init__(self, resources):
            self.resources = resources

        def get(self, _key):
            if isinstance(self.resources, Exception):
                raise self.resources
            return self.resources

    class Object:
        def __init__(self, value):
            self.value = value

        def get_object(self):
            return self.value

    assert worker._page_has_image(Page(None)) is False
    assert worker._page_has_image(Page(Object({}))) is False
    assert worker._page_has_image(Page(Object({"/XObject": Object({})}))) is False
    assert worker._page_has_image(Page(Object(
        {"/XObject": Object({"x": Object({"/Subtype": "/Image"})})}
    ))) is True
    assert worker._page_has_image(Page(RuntimeError("bad resources"))) is False


def test_parse_pdf_budget_all_failed_and_page_limit_paths(monkeypatch):
    limits = _profile()["limits"]
    identity = {"profile_id": "synthetic"}
    pages = _pdf([
        [(100, 750, "First page.")],
        [(100, 750, "Second page.")],
    ])
    zero_budget = worker._parse_pdf(pages, None, 0, limits, identity)
    assert zero_budget["pages"][0]["status"] == "truncated"
    assert zero_budget["truncated"] is True

    limited = worker._parse_pdf(pages, None, 200_000, {**limits, "pdf_pages": 1}, identity)
    assert limited["status"] == "truncated"
    assert "page_limit_truncated" in limited["warnings"]

    original = worker._pdf_fragments

    def always_fail(_page, _limits):
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(worker, "_pdf_fragments", always_fail)
    failed = worker._parse_pdf(pages, None, 200_000, limits, identity)
    assert failed["status"] == "failed" and failed["error"]["code"] == "parse_failed"
    monkeypatch.setattr(worker, "_pdf_fragments", original)


def test_parse_pdf_blank_and_scanned_pages_in_process():
    blank = worker._parse_pdf(
        _pdf([[]]), None, 200_000, _profile()["limits"], {"profile_id": "synthetic"}
    )
    assert blank["pages"][0]["status"] == "empty"

    from PIL import Image
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    image = Image.new("RGB", (40, 20), "white")
    encoded = io.BytesIO()
    image.save(encoded, "PNG")
    encoded.seek(0)
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawImage(ImageReader(encoded), 100, 600, width=40, height=20)
    document.showPage()
    document.save()
    scanned = worker._parse_pdf(
        buffer.getvalue(), None, 200_000, _profile()["limits"],
        {"profile_id": "synthetic"},
    )
    assert scanned["status"] == "partial"
    assert scanned["pages"][0]["status"] == "partial"
    assert scanned["vision_required"] is True


def test_preflight_docx_compression_xml_and_content_type_rejections():
    def archive(entries, compression=zipfile.ZIP_STORED):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression) as target:
            for name, data in entries.items():
                target.writestr(name, data)
        return buffer.getvalue()

    limits = _profile()["limits"]
    compressed = dict(limits)
    compressed["docx_compression_ratio"] = 1
    with pytest.raises(worker.WorkerFailure, match="archive_limit"):
        worker._preflight_docx(
            archive({"word/document.xml": b"x" * (1024 * 1024 + 1)},
                    zipfile.ZIP_DEFLATED),
            compressed,
        )
    with pytest.raises(worker.WorkerFailure, match="malformed_document"):
        worker._preflight_docx(
            archive({"word/document.xml": b"<broken"}), limits
        )
    with pytest.raises(worker.WorkerFailure, match="active_content_rejected"):
        worker._preflight_docx(
            archive({
                "word/document.xml": b"<x/>",
                "[Content_Types].xml": (
                    b'<Types><Override ContentType="application/vnd.ms-word.'
                    b'document.macroEnabled.main+xml"/></Types>'
                ),
            }),
            limits,
        )


def test_parse_docx_failure_limits_and_continuation():
    limits = _profile()["limits"]
    failed = worker._parse_docx(b"not-a-docx", 200_000, limits, {"profile_id": "x"})
    assert failed["status"] == "failed" and failed["error"]["code"] == "malformed_document"

    payload = _docx(["A paragraph that is long enough."],
                    [[["Item", "Value"], ["Row", "1.5"]]])
    block_limited = worker._parse_docx(
        payload, 200_000, {**limits, "blocks": 0}, {"profile_id": "x"}
    )
    assert block_limited["status"] == "truncated" and block_limited["blocks"] == []

    row_limited = worker._parse_docx(
        payload, 200_000, {**limits, "table_rows": 0}, {"profile_id": "x"}
    )
    assert row_limited["truncated"] is True

    cell_limited = worker._parse_docx(
        payload, 200_000, {**limits, "table_cells": 0}, {"profile_id": "x"}
    )
    assert cell_limited["truncated"] is True

    zero_budget = worker._parse_docx(payload, 0, limits, {"profile_id": "x"})
    assert zero_budget["status"] == "truncated" and zero_budget["text"] == ""

    short_budget = worker._parse_docx(payload, 3, limits, {"profile_id": "x"})
    assert short_budget["status"] == "truncated" and short_budget["text"] == "A p"

    continued = worker._parse_docx(
        _docx(tables=[
            [["Item", "Value"], ["First", "1.5"]],
            [["Item", "Value"], ["Second", "2.5"]],
        ]),
        200_000,
        limits,
        {"profile_id": "x"},
    )
    assert continued["tables"][1]["continuation_of"] == "body/tbl[0]"


def test_profile_float_dependency_shape_and_missing_package():
    profile = _profile()
    profile["limits"]["deadline_seconds"] = 0.0
    with pytest.raises(worker.WorkerFailure, match="invalid_profile"):
        worker._validate_profile(profile)
    profile = _profile()
    profile["dependencies"] = {}
    with pytest.raises(worker.WorkerFailure, match="invalid_profile"):
        worker._validate_profile(profile)
    profile = _profile()
    profile["dependencies"] = [{"name": "definitely-not-installed", "version": "1"}]
    with pytest.raises(worker.WorkerFailure, match="dependency_missing"):
        worker._verify_dependencies(profile)


def test_read_json_and_write_result_failure_boundaries(tmp_path):
    with pytest.raises(worker.WorkerFailure, match="request_too_large"):
        worker._read_json(_write(tmp_path / "large.json", b"x" * 10), 5)
    with pytest.raises(worker.WorkerFailure, match="invalid_request"):
        worker._read_json(_write(tmp_path / "bad.json", b"{"), 100)
    with pytest.raises(worker.WorkerFailure, match="invalid_request"):
        worker._read_json(_write(tmp_path / "list.json", b"[]"), 100)
    with pytest.raises(worker.WorkerFailure, match="io_failure"):
        worker._write_result(tmp_path, {"status": "failed"}, 100)


def _write(path, payload):
    path.write_bytes(payload)
    return path


def test_run_missing_profile_and_invalid_profile_json(tmp_path, monkeypatch):
    request, _output, profile_path = _request(
        tmp_path, monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    profile_path.unlink()
    with pytest.raises(worker.WorkerFailure, match="missing_profile"):
        worker._run(request)

    request, _output, profile_path = _request(
        tmp_path / "invalid", monkeypatch, _pdf([[(100, 750, "Text.")]])
    )
    profile_path.write_bytes(b"{")
    value = json.loads(request.read_text())
    value["profile_sha256"] = hashlib.sha256(b"{").hexdigest()
    request.write_text(json.dumps(value))
    with pytest.raises(worker.WorkerFailure, match="invalid_profile"):
        worker._run(request)


def test_main_generic_failure_and_network_guard_wiring(tmp_path, monkeypatch):
    request, _output, _profile_path = _request(
        tmp_path, monkeypatch, _pdf([[(100, 750, "Text.")]]), cancel_guard=False
    )

    def explode(_path):
        raise RuntimeError("synthetic unexpected failure")

    monkeypatch.setattr(worker, "_run", explode)
    assert worker.main(["worker", str(request)]) == 4

    import socket
    import urllib.request

    hooks = []
    monkeypatch.setattr(worker.sys, "addaudithook", hooks.append)
    original_create = socket.create_connection
    original_address = socket.getaddrinfo
    original_urlopen = urllib.request.urlopen
    try:
        worker._install_network_guard()
        assert hooks and socket.create_connection is worker._network_denied
        assert socket.getaddrinfo is worker._network_denied
        assert urllib.request.urlopen is worker._network_denied
    finally:
        socket.create_connection = original_create
        socket.getaddrinfo = original_address
        urllib.request.urlopen = original_urlopen
