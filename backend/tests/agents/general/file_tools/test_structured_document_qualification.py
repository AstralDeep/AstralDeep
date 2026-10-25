"""Synthetic qualification corpus comparing the baseline document reader with the opt-in
structured adapter on reading order, table association, provenance, and truthful status.
"""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

from agents.general.file_tools import read_document as read_document_module
from agents.general.file_tools import structured_worker
from agents.general.file_tools.read_document import _read_docx, _read_pdf
from agents.general.file_tools.structured_document import parse_structured_document


def _pdf(pages):
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    for operators in pages:
        document.setFont("Helvetica", 10)
        for x, y, text in operators:
            document.drawString(x, y, text)
        document.showPage()
    document.save()
    return buffer.getvalue()


def _table_pdf(rows):
    return _pdf([[
        (80 + 120 * column, 740 - 20 * row, value)
        for row, values in enumerate(rows)
        for column, value in enumerate(values)
    ]])


def _columns_pdf():
    rows = [
        ("Left column first sentence with enough words", "Right column first sentence with enough words"),
        ("Left column second sentence with enough words", "Right column second sentence with enough words"),
        ("Left column third sentence with enough words", "Right column third sentence with enough words"),
    ]
    return _pdf([[
        item
        for index, (left, right) in enumerate(rows)
        for item in ((60, 740 - 18 * index, left), (330, 740 - 18 * index, right))
    ]])


def _scanned_pdf():
    from PIL import Image
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    image = Image.new("RGB", (160, 80), "white")
    encoded = io.BytesIO()
    image.save(encoded, "PNG")
    encoded.seek(0)
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawImage(ImageReader(encoded), 100, 600, width=160, height=80)
    document.showPage()
    document.save()
    return buffer.getvalue()


def _docx_table():
    import docx

    document = docx.Document()
    document.add_paragraph("Clinical table follows.")
    table = document.add_table(rows=2, cols=3)
    for column, value in enumerate(["Analyte", "Value", "Unit"]):
        table.rows[0].cells[column].text = value
    for column, value in enumerate(["Sodium", "-1.25", "mg"]):
        table.rows[1].cells[column].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_continuation():
    import docx

    document = docx.Document()
    for index in range(2):
        table = document.add_table(rows=2, cols=2)
        for column, value in enumerate(["Item", "Value"]):
            table.rows[0].cells[column].text = value
        for column, value in enumerate([f"Row {index}", "-2.50"]):
            table.rows[1].cells[column].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _base_status(result):
    return result.get("status", "failed" if "error" in result else "complete")


def test_synthetic_corpus_compares_baseline_transcription_and_candidate_structure(monkeypatch):
    monkeypatch.setattr(read_document_module, "pdf_to_vision_images", lambda _payload: [])

    columns = _columns_pdf()
    table = _table_pdf([
        ["Analyte", "Value", "Unit"],
        ["Sodium", "-1.25", "mg"],
        ["Glucose", "+2.50", "mmol/L"],
    ])
    scanned = _scanned_pdf()
    blank = _pdf([[]])
    malformed = b"%PDF-1.4\nbroken body \x00"
    docx_table = _docx_table()
    docx_continuation = _docx_continuation()

    baseline_columns = _read_pdf(columns, None, 200_000)
    candidate_columns = parse_structured_document(columns, "pdf", None, 200_000)
    assert baseline_columns["text"].index("Right column first") < baseline_columns["text"].index(
        "Left column third"
    )
    assert candidate_columns["text"].index("Left column third") < candidate_columns["text"].index(
        "Right column first"
    )

    baseline_table = _read_pdf(table, None, 200_000)
    candidate_table = parse_structured_document(table, "pdf", None, 200_000)
    assert "tables" not in baseline_table
    assert candidate_table["tables"][0]["rows"][1]["cells"][1]["text"] == "-1.25"
    assert candidate_table["tables"][0]["rows"][2]["cells"][1]["text"] == "+2.50"
    assert candidate_table["tables"][0]["rows"][2]["cells"][2]["text"] == "mmol/L"

    baseline_scanned = _read_pdf(scanned, None, 200_000)
    candidate_scanned = parse_structured_document(scanned, "pdf", None, 200_000)
    assert baseline_scanned["vision_required"] is True
    assert candidate_scanned["status"] == "partial"
    assert candidate_scanned["vision_required"] is True
    assert candidate_scanned["text"] == ""

    baseline_blank = _read_pdf(blank, None, 200_000)
    candidate_blank = parse_structured_document(blank, "pdf", None, 200_000)
    assert baseline_blank["vision_required"] is True
    assert candidate_blank["status"] == "complete" and candidate_blank["text"] == ""

    candidate_malformed = parse_structured_document(malformed, "pdf", None, 200_000)
    try:
        _read_pdf(malformed, None, 200_000)
    except Exception:
        baseline_malformed_status = "failed"
    else:
        baseline_malformed_status = "complete"
    assert candidate_malformed["status"] == "failed"

    baseline_docx = _read_docx(docx_table, 200_000)
    candidate_docx = parse_structured_document(docx_table, "docx", None, 200_000)
    assert "Sodium" not in baseline_docx["text"]
    assert candidate_docx["tables"][0]["rows"][1]["cells"][0]["text"] == "Sodium"
    assert candidate_docx["tables"][0]["rows"][1]["cells"][1]["text"] == "-1.25"

    candidate_continuation = parse_structured_document(
        docx_continuation, "docx", None, 200_000
    )
    assert candidate_continuation["tables"][1]["continuation_of"] == "body/tbl[0]"

    statuses = {
        "columns": _base_status(candidate_columns),
        "table": _base_status(candidate_table),
        "scanned": _base_status(candidate_scanned),
        "blank": _base_status(candidate_blank),
        "malformed": _base_status(candidate_malformed),
        "docx_table": _base_status(candidate_docx),
        "docx_continuation": _base_status(candidate_continuation),
    }
    assert statuses == {
        "columns": "complete",
        "table": "complete",
        "scanned": "partial",
        "blank": "complete",
        "malformed": "failed",
        "docx_table": "complete",
        "docx_continuation": "complete",
    }
    assert list(statuses.values()).count("failed") == 1
    assert list(statuses.values()).count("partial") == 1
    assert baseline_malformed_status == "failed"
    assert candidate_columns["source"]["sha256"] == hashlib.sha256(columns).hexdigest()


def test_qualification_provenance_is_page_and_region_bounded():
    payload = _table_pdf([["Item", "Value"], ["Sodium", "-1.25"]])
    result = parse_structured_document(payload, "pdf", None, 200_000)
    assert result["page_count"] == 1 and result["selected_pages"] == [1]
    for block in result["blocks"]:
        source = block["source"]
        assert source["page"] in result["selected_pages"]
        region = source["region"]
        assert len(region) == 4
        assert 0 <= region[0] < region[2]
        assert 0 <= region[1] < region[3]
    for table in result["tables"]:
        assert table["source"]["page"] == 1
        for row in table["rows"]:
            for cell in row["cells"]:
                assert cell["source"]["region"][0] < cell["source"]["region"][2]


def test_pdf_partial_page_failure_is_reported_not_promoted(monkeypatch):
    payload = _pdf([
        [(100, 750, "First page remains readable.")],
        [(100, 750, "Second page fails during extraction.")],
    ])
    profile = json.loads(
        (Path(__file__).resolve().parents[5] / "config" / "documents" /
         "structured-reader-profile.json").read_text()
    )
    original = structured_worker._pdf_fragments
    calls = {"count": 0}

    def fail_second(page, limits):
        calls["count"] += 1
        if calls["count"] == 2:
            raise RuntimeError("synthetic page failure")
        return original(page, limits)

    monkeypatch.setattr(structured_worker, "_pdf_fragments", fail_second)
    result = structured_worker._parse_pdf(
        payload, None, 200_000, profile["limits"], {"profile_id": "synthetic"}
    )
    assert result["status"] == "partial" and result["complete"] is False
    assert [page["status"] for page in result["pages"]] == ["complete", "failed"]
    assert "page_extraction_failed" in result["warnings"]
    assert "First page remains readable." in result["text"]
    assert "Second page fails" not in result["text"]